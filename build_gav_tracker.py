#!/usr/bin/env python3
"""
Pipeline: F01CGPPTO -> baseline JSON -> seguimiento_gavs.html

Lee el archivo de presupuesto de GAVs/Indirectos de Exportaciones y genera el
baseline que consume el tracker HTML:

  * hoja `Formato`  -> montos autoritativos por línea y por mes (todo el ppto)
  * hoja `BD_Opex`  -> partidas con Denominación y Texto de cabecera para los
                       recursos en foco (el desglose fino que se va a asignar)

El control se concentra en los recursos de FOCO; el resto del presupuesto queda
como referencia de solo lectura.

Uso:
    python build_gav_tracker.py /ruta/al/F01CGPPTO_..._Consolidado.xlsx
"""

import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci
from openpyxl.utils import get_column_letter as gcl

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "gav_template.html"
OUT_HTML = HERE / "seguimiento_gavs.html"
OUT_JSON = HERE / "gav_baseline.json"

# --- recursos bajo control activo -------------------------------------------
FOCO = ["Gastos De Viaje", "Gastos de Feria y Eventos", "Refrigerios"]

# --- temporada ---------------------------------------------------------------
MONTHS = [f"2026_{m:02d}" for m in range(7, 13)] + [f"2027_{m:02d}" for m in range(1, 7)]
QTY_FIRST = "AD"   # (Q)2026_07 en la hoja Formato
USD_FIRST = "BN"   # (M)2026_07 en USD -> la fila 13 trae el TC presupuestado
SEASON = {"id": "T2627", "label": "Temporada 26/27"}

HEADER_ROW, TC_ROW, FIRST_DATA_ROW = 15, 13, 16
TC_FALLBACK = 3.36
NA = re.compile(r"#N/?A", re.I)


def val(row, letter):
    v = row[ci(letter) - 1]
    if isinstance(v, str):
        return v.strip() or None
    return v


def col_at(first, offset):
    return gcl(ci(first) + offset)


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


def txt(v):
    s = "" if v is None else str(v).strip()
    return "" if s.lower() in ("none", "nan", "#n/a") else s


# ---------------------------------------------------------------- hoja Formato
def read_formato(wb):
    ws = wb["Formato"]
    head = list(ws.iter_rows(min_row=1, max_row=5, max_col=2, values_only=True))
    header = {
        "entidad": txt(head[0][1]),
        "presupuesto": txt(head[1][1]),
        "ambito": txt(head[2][1]),
        "gerencia": txt(head[3][1]),
    }

    rows = list(
        ws.iter_rows(min_row=TC_ROW, max_row=ws.max_row, max_col=ci("BZ"), values_only=True)
    )
    tc = [num(val(rows[0], col_at(USD_FIRST, i))) or TC_FALLBACK for i in range(len(MONTHS))]

    lines = []
    for r in rows[FIRST_DATA_ROW - TC_ROW :]:
        if not val(r, "E"):
            continue
        qty = [num(val(r, col_at(QTY_FIRST, i))) for i in range(len(MONTHS))]
        if not any(qty):
            continue

        recurso = val(r, "H")
        if not recurso or NA.search(str(recurso)):
            # el VLOOKUP del formato falla para algunas cuentas; AgrCtas lo resuelve
            recurso = val(r, "G") or "Sin clasificar"

        pu, moneda = num(val(r, "U")), (val(r, "X") or "USD")
        usd = [(q * pu / tc[i]) if moneda == "PEN" else q * pu for i, q in enumerate(qty)]

        lines.append(
            {
                "id": f"L{len(lines)+1:03d}",
                "ceco": val(r, "E"),
                "recurso": recurso,
                "claseCoste": txt(val(r, "J")),
                "denomClaseCoste": txt(val(r, "K")),
                "descripClaseCoste": txt(val(r, "L")),
                "material": txt(val(r, "M")),
                "textoMaterial": txt(val(r, "N")),
                "tipo": txt(val(r, "Q")),
                "umb": txt(val(r, "R")),
                "pu": round(pu, 6),
                "moneda": moneda,
                "usd": [round(x, 2) for x in usd],
            }
        )
    return header, [round(x, 4) for x in tc], lines


# ---------------------------------------------------------------- hoja BD_Opex
def read_bdopex(wb, den2rec, cat):
    """Partidas con Denominación y Texto de cabecera, filtradas a la temporada.

    `Denom.clase de coste` llega truncada desde SAP ("Gastos de Feria",
    "Consumos del Persona"), así que la clase de coste se normaliza contra el
    catálogo del Formato: manda `Descrip.clases coste`.
    """
    ws = wb["BD_Opex"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [txt(c) for c in rows[0]]
    ix = {n: i for i, n in enumerate(hdr)}

    need = ["Centro de coste", "Ejercicio", "Período", "Denom.clase de coste", "Val/Mon.so.CO"]
    missing = [c for c in need if c not in ix]
    if missing:
        print(f"[aviso] BD_Opex sin columnas {missing}; no se generan partidas")
        return []

    get = lambda r, c: txt(r[ix[c]]) if c in ix else ""
    parts = []
    for r in rows[1:]:
        if not r[ix["Centro de coste"]]:
            continue
        try:
            mes = f"{int(r[ix['Ejercicio']])}_{int(r[ix['Período']]):02d}"
        except (TypeError, ValueError):
            continue
        if mes not in MONTHS:
            continue
        den = get(r, "Denom.clase de coste")
        rec = den2rec.get(den)
        if rec not in FOCO:
            continue
        monto = num(r[ix["Val/Mon.so.CO"]])
        if abs(monto) < 0.005:
            continue
        cc = get(r, "Clase de coste")
        ref = cat.get(cc.strip()) or {}
        parts.append(
            {
                "id": f"P{len(parts)+1:03d}",
                "recurso": rec,
                "mi": MONTHS.index(mes),
                "claseCoste": cc,
                "denomClaseCoste": ref.get("denomClaseCoste") or den,
                "descripClaseCoste": ref.get("descripClaseCoste")
                or get(r, "Descrip.clases coste")
                or den,
                "denominacion": get(r, "Denominación"),
                "texto": get(r, "Texto de cabecera de documento"),
                "monto": round(monto, 2),
            }
        )
    return parts


def extract(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    header, tc, lines = read_formato(wb)

    den2rec = {}
    for l in lines:
        den2rec.setdefault(l["denomClaseCoste"], l["recurso"])

    cat = {}
    for l in lines:
        cat.setdefault(
            str(l["claseCoste"]).strip(),
            {
                "denomClaseCoste": l["denomClaseCoste"],
                "descripClaseCoste": l["descripClaseCoste"] or l["denomClaseCoste"],
            },
        )

    parts = read_bdopex(wb, den2rec, cat)

    # tope por recurso = el Formato manda
    topes = {r: 0.0 for r in FOCO}
    for l in lines:
        if l["recurso"] in topes:
            topes[l["recurso"]] += sum(l["usd"])
    topes = {k: round(v, 2) for k, v in topes.items()}

    # lo que el detalle de SAP no cubre queda como partida libre "Por asignar"
    for rec in FOCO:
        cubierto = sum(p["monto"] for p in parts if p["recurso"] == rec)
        resto = round(topes[rec] - cubierto, 2)
        if abs(resto) < 0.5:
            continue
        cand = [l for l in lines if l["recurso"] == rec]
        parts.append(
            {
                "id": f"P{len(parts)+1:03d}",
                "recurso": rec,
                "mi": 0,
                "claseCoste": cand[0]["claseCoste"] if cand else "",
                "denomClaseCoste": cand[0]["denomClaseCoste"] if cand else "",
                "descripClaseCoste": cand[0]["descripClaseCoste"] if cand else "",
                "denominacion": "",
                "texto": "Por asignar",
                "monto": resto,
            }
        )

    # catálogo de clases de coste por recurso, para los selectores de la app
    clases = {}
    for rec in FOCO:
        seen, out = set(), []
        for l in lines:
            if l["recurso"] != rec:
                continue
            k = (l["claseCoste"], l["denomClaseCoste"], l["descripClaseCoste"])
            if k in seen:
                continue
            seen.add(k)
            out.append(
                {
                    "cc": k[0],
                    "denomClaseCoste": k[1],
                    "descripClaseCoste": k[2] or k[1],  # la descripción es la etiqueta
                }
            )
        clases[rec] = out

    return {
        "header": header,
        "season": {**SEASON, "months": MONTHS, "tc": tc},
        "foco": FOCO,
        "topes": topes,
        "clases": clases,
        "parts": parts,
        "lines": lines,
        "totalFoco": round(sum(topes.values()), 2),
        "totalPpto": round(sum(sum(l["usd"]) for l in lines), 2),
    }


def build(xlsx_path):
    p = extract(xlsx_path)
    OUT_JSON.write_text(json.dumps(p, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"{p['season']['label']}: {len(p['lines'])} líneas · USD {p['totalPpto']:,.2f}")
    print(f"En foco ({len(FOCO)} recursos): USD {p['totalFoco']:,.2f} · {len(p['parts'])} partidas")
    for r in FOCO:
        n = sum(1 for x in p["parts"] if x["recurso"] == r)
        print(f"   {r:28} tope {p['topes'][r]:>9,.2f}  ·  {n:2} partidas")

    if TEMPLATE.exists():
        html = TEMPLATE.read_text(encoding="utf-8")
        html = html.replace("/*__BASELINE__*/", json.dumps(p, ensure_ascii=False))
        OUT_HTML.write_text(html, encoding="utf-8")
        print(f"\n{OUT_HTML.name} generado ({OUT_HTML.stat().st_size/1024:,.0f} KB)")
    else:
        print(f"\n[aviso] falta {TEMPLATE.name}; solo se generó {OUT_JSON.name}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    build(sys.argv[1])
